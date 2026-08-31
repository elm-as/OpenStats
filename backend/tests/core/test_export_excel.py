"""
Tests unitaires pour l'export Excel stylisé et les formules dynamiques recalculables.
"""

import os
import openpyxl
import pandas as pd
import pytest
from app.core.export_excel import export_excel


def test_export_excel_with_dynamic_formulas(tmp_path):
    """Vérifie que l'export Excel crée les feuilles et injecte les formules dynamiques."""
    output_file = str(tmp_path / "test_report.xlsx")

    payload = {
        "metadata": {"title": "Test Report", "organization": "Elmas Labs"},
        "dataset": {"id": "ds-1", "name": "Test Data", "original_filename": "test.csv"},
        "data_summary": {
            "preview": [
                {"prix": 100.0, "quantite": 10},
                {"prix": 120.0, "quantite": 15},
                {"prix": 110.0, "quantite": 12},
                {"prix": 130.0, "quantite": 20},
            ]
        },
        "analysis": {
            "descriptive_stats": {
                "prix": {
                    "type": "numeric",
                    "dtype": "float64",
                    "count": 4,
                    "mean": 115.0,
                    "median": 115.0,
                    "std": 12.9,
                    "min": 100.0,
                    "max": 130.0,
                },
                "quantite": {
                    "type": "numeric",
                    "dtype": "int64",
                    "count": 4,
                    "mean": 14.25,
                    "median": 13.5,
                    "std": 4.3,
                    "min": 10,
                    "max": 20,
                },
            }
        },
    }

    result_path = export_excel(output_file, payload)
    assert os.path.exists(result_path)

    # Charger le classeur pour inspecter les formules
    wb = openpyxl.load_workbook(result_path, data_only=False)
    assert "Apercu" in wb.sheetnames
    assert "Stats Numeriques" in wb.sheetnames

    ws_stats = wb["Stats Numeriques"]
    # Vérifier que la cellule Moyenne contient bien une formule Excel dynamique
    # Trouver l'indice de la colonne Moyenne
    moyenne_col = None
    for col in range(1, ws_stats.max_column + 1):
        if ws_stats.cell(row=1, column=col).value == "Moyenne":
            moyenne_col = col
            break

    assert moyenne_col is not None
    cell_val = str(ws_stats.cell(row=2, column=moyenne_col).value)
    assert cell_val.startswith("=AVERAGE(")
    assert "'Apercu'!" in cell_val
