"""
Module d'export Excel multi-onglets stylisé.
"""

from __future__ import annotations

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.core.export_formatters import _safe_sheet_name
from app.core.export_frames_stats import (
    _summary_frame, _preview_frame, _dictionary_frame, _descriptive_frames,
    _correlation_matrix_frame, _significant_correlations_frame, _vif_frame,
    _tests_frame, _log_frame, _versions_frame, _history_frame, _audit_frame
)
from app.core.export_frames_models import (
    _model_ranking_frame, _feature_importance_frame, _shap_frame,
    _timeseries_summary_frame, _timeseries_forecast_frame,
    _multivariate_summary_frame, _granger_frame, _johansen_frame,
    _multivariate_forecast_frame, _pca_variance_frame, _pca_loadings_frame,
    _ca_coords_frame, _mca_modalities_frame
)


def _inject_dynamic_excel_formulas(workbook) -> None:
    """Injecte de vraies formules Excel recalculables reliant Stats Numeriques à Apercu."""
    if "Apercu" not in workbook.sheetnames or "Stats Numeriques" not in workbook.sheetnames:
        return

    ws_data = workbook["Apercu"]
    ws_stats = workbook["Stats Numeriques"]
    max_data_row = ws_data.max_row
    if max_data_row < 2:
        return

    # Index des colonnes de données dans Apercu
    data_cols: dict[str, str] = {}
    for col_idx in range(1, ws_data.max_column + 1):
        val = str(ws_data.cell(row=1, column=col_idx).value or "").strip()
        if val:
            data_cols[val] = get_column_letter(col_idx)

    # Index des colonnes cibles dans Stats Numeriques
    stats_cols: dict[str, int] = {}
    for col_idx in range(1, ws_stats.max_column + 1):
        val = str(ws_stats.cell(row=1, column=col_idx).value or "").strip()
        if val:
            stats_cols[val] = col_idx

    var_col_idx = stats_cols.get("Variable")
    if not var_col_idx:
        return

    for row_idx in range(2, ws_stats.max_row + 1):
        var_name = str(ws_stats.cell(row=row_idx, column=var_col_idx).value or "").strip()
        col_letter = data_cols.get(var_name)
        if not col_letter:
            continue

        rng = f"'Apercu'!${col_letter}$2:${col_letter}${max_data_row}"

        if "Effectif" in stats_cols:
            ws_stats.cell(row=row_idx, column=stats_cols["Effectif"], value=f"=COUNT({rng})")
        if "Moyenne" in stats_cols:
            ws_stats.cell(row=row_idx, column=stats_cols["Moyenne"], value=f"=AVERAGE({rng})")
        if "Mediane" in stats_cols:
            ws_stats.cell(row=row_idx, column=stats_cols["Mediane"], value=f"=MEDIAN({rng})")
        if "Ecart-type" in stats_cols:
            ws_stats.cell(row=row_idx, column=stats_cols["Ecart-type"], value=f"=STDEV.S({rng})")
        if "Variance" in stats_cols:
            ws_stats.cell(row=row_idx, column=stats_cols["Variance"], value=f"=VAR.S({rng})")
        if "Min" in stats_cols:
            ws_stats.cell(row=row_idx, column=stats_cols["Min"], value=f"=MIN({rng})")
        if "Max" in stats_cols:
            ws_stats.cell(row=row_idx, column=stats_cols["Max"], value=f"=MAX({rng})")


def export_excel(output_path: str, payload: dict) -> str:
    """Exporte les resultats dans un classeur Excel multi-onglets."""
    numeric_stats, categorical_stats = _descriptive_frames(payload)
    used_sheet_names: set[str] = set()
    tables = [
        ("Resume", _summary_frame(payload)),
        ("Apercu", _preview_frame(payload)),
        ("Dictionnaire", _dictionary_frame(payload)),
        ("Stats Numeriques", numeric_stats),
        ("Stats Categorielles", categorical_stats),
        ("Corr Pearson", _correlation_matrix_frame(payload, "pearson")),
        ("Corr Spearman", _correlation_matrix_frame(payload, "spearman")),
        ("Corr Significatives", _significant_correlations_frame(payload)),
        ("VIF", _vif_frame(payload)),
        ("Tests", _tests_frame(payload)),
        ("Modelisation", _model_ranking_frame(payload)),
        ("Importance Vars", _feature_importance_frame(payload)),
        ("SHAP Global", _shap_frame(payload)),
        ("TS Resume", _timeseries_summary_frame(payload)),
        ("TS Previsions", _timeseries_forecast_frame(payload)),
        ("MTS Resume", _multivariate_summary_frame(payload)),
        ("Granger", _granger_frame(payload)),
        ("Johansen", _johansen_frame(payload)),
        ("MTS Previsions", _multivariate_forecast_frame(payload)),
        ("ACP Variance", _pca_variance_frame(payload)),
        ("ACP Loadings", _pca_loadings_frame(payload)),
        ("AFC Lignes", _ca_coords_frame(payload, axis="row")),
        ("AFC Colonnes", _ca_coords_frame(payload, axis="col")),
        ("ACM Modalites", _mca_modalities_frame(payload)),
        ("Nettoyage", _log_frame(payload.get("cleaning_log"))),
        ("Transformations", _log_frame(payload.get("transform_logs"))),
        ("Versions", _versions_frame(payload)),
        ("Historique", _history_frame(payload)),
        ("Audit", _audit_frame(payload)),
    ]

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        for base_name, frame in tables:
            if frame is None or frame.empty:
                continue
            sheet_name = _safe_sheet_name(base_name, used_sheet_names)
            frame.to_excel(writer, sheet_name=sheet_name, index=False)

        workbook = writer.book
        _inject_dynamic_excel_formulas(workbook)

        header_fill = PatternFill(start_color="16324F", end_color="16324F", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True, size=11)
        thin_border = Border(
            left=Side(style="thin", color="D5DCE5"),
            right=Side(style="thin", color="D5DCE5"),
            top=Side(style="thin", color="D5DCE5"),
            bottom=Side(style="thin", color="D5DCE5"),
        )

        for worksheet in workbook.worksheets:
            worksheet.freeze_panes = "A2"
            worksheet.sheet_view.showGridLines = False
            if worksheet.max_row >= 1:
                for cell in worksheet[1]:
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                    cell.border = thin_border
            if worksheet.max_row > 1:
                worksheet.auto_filter.ref = worksheet.dimensions
            for row in worksheet.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
                    cell.border = thin_border
            for col_idx, col_cells in enumerate(worksheet.columns, 1):
                max_len = max((len(str(cell.value or "")) for cell in col_cells), default=10)
                worksheet.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 3, 12), 60)

    return output_path
